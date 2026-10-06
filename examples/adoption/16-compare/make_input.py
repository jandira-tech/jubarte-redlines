# SPDX-FileCopyrightText: 2026 Jandira Technologies, LLC
# SPDX-License-Identifier: AGPL-3.0-only
"""Build v1.docx and v2.docx: two drafts of one letter.

v2 differs from v1 by:
  - two word changes ("2 February 2026" -> "9 February 2026",
    "thirty days" -> "fifteen days"),
  - one paragraph moved (the work-product clause moves from after the
    termination clause to right after the scope clause),
  - one paragraph deleted (the termination clause).
"""

import datetime as _dt

from docx import Document

V1 = [
    ("h", "Services Agreement (Draft)"),
    ("p", "This agreement is between Jandira Technologies, LLC and Northwind Traders Ltd."),
    ("p", "Jandira will deliver the invoicing redesign described in Schedule A."),
    ("p", "Work begins on 2 February 2026 and ends on 30 June 2026."),
    ("p", "Invoices are payable within thirty days of receipt."),
    ("p", "Either party may terminate with two weeks' written notice."),
    ("p", "All work product is assigned to Northwind on payment."),
    ("p", "Signed in duplicate on 12 January 2026."),
]

V2 = [
    ("h", "Services Agreement (Draft)"),
    ("p", "This agreement is between Jandira Technologies, LLC and Northwind Traders Ltd."),
    ("p", "Jandira will deliver the invoicing redesign described in Schedule A."),
    ("p", "All work product is assigned to Northwind on payment."),  # moved up
    ("p", "Work begins on 9 February 2026 and ends on 30 June 2026."),  # word change
    ("p", "Invoices are payable within fifteen days of receipt."),  # word change
    ("p", "Signed in duplicate on 12 January 2026."),
    # "Either party may terminate..." deleted
]


def build(paras, path):
    doc = Document()
    for kind, text in paras:
        if kind == "h":
            doc.add_heading(text, level=1)
        else:
            doc.add_paragraph(text)
    props = doc.core_properties
    props.created = _dt.datetime(2026, 1, 12, 9, 0, tzinfo=_dt.timezone.utc)
    props.modified = _dt.datetime(2026, 1, 12, 9, 0, tzinfo=_dt.timezone.utc)
    props.last_modified_by = "People Operations"
    props.author = "People Operations"
    doc.save(path)
    print(f"wrote {path}")


build(V1, "v1.docx")
build(V2, "v2.docx")
