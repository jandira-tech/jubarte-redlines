#!/usr/bin/env python3
# SPDX-FileCopyrightText: 2026 Jandira Technologies, LLC
#
# SPDX-License-Identifier: AGPL-3.0-only

"""Write input.docx, the file every reader in this folder consumes.

python-docx builds the body (headings, bold/italic runs, a list, a table)
plus the header and footer, which a Markdown writer cannot produce. The
footnote is added afterwards with the standard library's zipfile: python-docx
1.2.0 has no footnote API, so footnotes.xml, its relationship, its content
type and the reference run are injected into the saved package directly.

    python3 make_input.py            # writes input.docx
"""

from __future__ import annotations

import shutil
import zipfile
from pathlib import Path

import docx

OUT = Path(__file__).resolve().parent / "input.docx"

FOOTNOTE_PART = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    '<w:footnote w:type="separator" w:id="-1"><w:p><w:pPr>'
    '<w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
    "<w:r><w:separator/></w:r></w:p></w:footnote>"
    '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:pPr>'
    '<w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
    "<w:r><w:continuationSeparator/></w:r></w:p></w:footnote>"
    '<w:footnote w:id="2"><w:p><w:pPr><w:pStyle w:val="FootnoteText"/></w:pPr>'
    '<w:r><w:rPr><w:rStyle w:val="FootnoteReference"/></w:rPr><w:footnoteRef/></w:r>'
    '<w:r><w:t xml:space="preserve"> Net thirty (30) days from the invoice date.</w:t></w:r>'
    "</w:p></w:footnote></w:footnotes>"
)

ANCHOR = "net thirty (30) days"

FOOTNOTE_REF_RUN = (
    '<w:r><w:rPr><w:rStyle w:val="FootnoteReference"/></w:rPr>'
    '<w:footnoteReference w:id="2"/></w:r>'
)


def build() -> None:
    doc = docx.Document()

    header = doc.sections[0].header
    header.paragraphs[0].text = "Jandira Technologies, LLC - Professional Services Agreement (Draft 3)"
    footer = doc.sections[0].footer
    footer.paragraphs[0].text = "Confidential draft - do not circulate"

    doc.add_heading("Professional Services Agreement", level=1)
    doc.add_paragraph(
        "This Agreement is made between Jandira Technologies, LLC ("
        '"the Company") and Northwind Traders Ltd. ("the Client"), '
        "effective as of 2 March 2026."
    )

    doc.add_heading("1. Scope of Services", level=2)
    p = doc.add_paragraph(style="Normal")
    r = p.add_run("Scope. ")
    r.bold = True
    p.add_run(
        "The Company will design, implement and document the document-generation "
        "pipeline described in Statement of Work 1, including acceptance tests."
    )

    doc.add_paragraph("The deliverables are:", style="Normal")
    for item in (
        "a written architecture summary",
        "the pipeline source and its tests",
        "a runbook for the operations team",
        "one training session of up to three hours",
    ):
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("2. Fees and Payment", level=2)
    p = doc.add_paragraph(style="Normal")
    r = p.add_run("Fees. ")
    r.bold = True
    r2 = p.add_run("Invoices are payable ")
    r2.italic = True
    p.add_run(ANCHOR + ", by wire transfer to the account stated on each invoice.")
    doc.add_paragraph(
        "An invoice that is not disputed within ten business days is deemed accepted.",
        style="Normal",
    )

    doc.add_paragraph("Payment schedule:", style="Normal")
    table = doc.add_table(rows=4, cols=3, style="Table Grid")
    head = ("Milestone", "Deliverable", "Fee (USD)")
    for i, text in enumerate(head):
        run = table.rows[0].cells[i].paragraphs[0].add_run(text)
        run.bold = True
    rows = (
        ("M1", "Architecture summary accepted", "12,000"),
        ("M2", "Pipeline passes acceptance tests", "30,000"),
        ("M3", "Runbook and training delivered", "8,000"),
    )
    for r_i, row in enumerate(rows, start=1):
        for c_i, text in enumerate(row):
            table.rows[r_i].cells[c_i].text = text

    doc.add_heading("3. Confidentiality", level=2)
    doc.add_paragraph(
        "Each party keeps the other's non-public information secret and uses it "
        "only to perform this Agreement. The obligation survives termination by "
        "three years.",
        style="Normal",
    )

    doc.add_heading("4. Termination", level=2)
    doc.add_paragraph(
        "Either party may terminate for material breach on fifteen days' written "
        "notice if the breach remains uncured. The Company is paid for work "
        "performed up to the termination date.",
        style="Normal",
    )

    doc.add_paragraph("Signed for the Company: ______________________", style="Normal")
    doc.add_paragraph("Signed for the Client:  ______________________", style="Normal")

    doc.save(OUT)


def inject_footnote() -> None:
    """Add word/footnotes.xml and the reference run to the saved package."""
    tmp = OUT.with_suffix(".tmp.docx")
    with zipfile.ZipFile(OUT) as zin:
        names = zin.namelist()
        data = {n: zin.read(n) for n in names}

    content_types = data["[Content_Types].xml"].decode("utf-8")
    assert "footnotes+xml" not in content_types
    content_types = content_types.replace(
        "</Types>",
        '<Override PartName="/word/footnotes.xml" ContentType="application/vnd'
        '.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>',
    )
    data["[Content_Types].xml"] = content_types.encode("utf-8")

    rels = data["word/_rels/document.xml.rels"].decode("utf-8")
    rels = rels.replace(
        "</Relationships>",
        '<Relationship Id="rIdFootnotes1" Type="http://schemas.openxmlformats.org'
        "/officeDocument/2006/relationships/footnotes\" "
        'Target="footnotes.xml"/></Relationships>',
    )
    data["word/_rels/document.xml.rels"] = rels.encode("utf-8")

    body = data["word/document.xml"].decode("utf-8")
    at = body.index(ANCHOR) + len(ANCHOR)
    end_run = body.index("</w:r>", at) + len("</w:r>")
    body = body[:end_run] + FOOTNOTE_REF_RUN + body[end_run:]
    data["word/document.xml"] = body.encode("utf-8")

    data["word/footnotes.xml"] = FOOTNOTE_PART.encode("utf-8")

    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in names:  # keep the original member order
            zout.writestr(n, data[n])
        zout.writestr("word/footnotes.xml", data["word/footnotes.xml"])
    shutil.move(tmp, OUT)


if __name__ == "__main__":
    build()
    inject_footnote()
    print(f"wrote {OUT}")
