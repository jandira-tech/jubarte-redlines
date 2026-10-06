# SPDX-FileCopyrightText: 2026 Jandira Technologies, LLC
# SPDX-License-Identifier: AGPL-3.0-only
"""The python-docx side of 09-edit-tracked.

Task: replace one sentence and delete one paragraph in input.docx.

python-docx has no tracked-changes API: it cannot write w:ins or w:del, so
whatever it saves looks like the original text. The sentence replacement
uses the common "rebuild the paragraph text" pattern; the paragraph
deletion has to drop the underlying XML element by hand.
"""
import docx

OLD = "within thirty days of receipt"
NEW = "within fifteen days of receipt"
DROP_PREFIX = "This Agreement begins on January 1, 2026"

doc = docx.Document("input.docx")

replaced = deleted = 0
for p in list(doc.paragraphs):
    if OLD in p.text:
        p.text = p.text.replace(OLD, NEW)
        replaced += 1

for p in list(doc.paragraphs):
    if p.text.startswith(DROP_PREFIX):
        p._element.getparent().remove(p._element)
        deleted += 1

doc.save("edit_pydocx.docx")
print(f"python-docx: {replaced} sentence replaced, {deleted} paragraph deleted")
print("python-docx: wrote edit_pydocx.docx with no w:ins/w:del (untracked edits)")
