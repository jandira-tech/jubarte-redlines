# SPDX-FileCopyrightText: 2026 Jandira Technologies, LLC
# SPDX-License-Identifier: AGPL-3.0-only
"""The python-docx side of 10-existing-revisions-keep.

Task: edit received.docx - a redline authored by "Counterparty" carrying
four tracked changes - the way an agent routinely does: find a paragraph
by text and rewrite it, while Counterparty's pending w:ins/w:del are
supposed to survive.

python-docx does not model revisions. paragraph.text is built from the
direct w:r children only, so it shows neither their w:del text nor their
w:ins text; and assigning paragraph.text replaces the paragraph's runs,
taking the w:ins/w:del elements with them.
"""
import docx

doc = docx.Document("received.docx")

for p in doc.paragraphs:
    if p.text.startswith("Client shall pay Contractor the fees"):
        print("fees paragraph as python-docx sees it:")
        print("  " + repr(p.text))
        print("  their inserted 'thirty' visible in p.text:", "thirty" in p.text)
        print("  their deleted 'forty-five' visible in p.text:", "forty-five" in p.text)

edits = 0
for p in doc.paragraphs:
    if "ninety days written notice" in p.text:
        p.text = p.text.replace("ninety days", "sixty days")
        edits += 1
    if p.text.startswith("Client shall pay Contractor the fees"):
        p.text = p.text.replace(
            "percent per month.", "percent per month, compounded daily.")
        edits += 1

doc.save("keep_pydocx.docx")
print(f"python-docx: {edits} edits applied -> keep_pydocx.docx")
