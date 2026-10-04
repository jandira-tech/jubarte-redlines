# SPDX-FileCopyrightText: 2026 Jandira Technologies, LLC
# SPDX-License-Identifier: AGPL-3.0-only
"""breaks.docx: a page break inside a paragraph, a page break in a
paragraph of its own, and a second section. LibreOffice's "MS Word 97"
export of it is breaks.doc:

    python3 make_breaks.py
    soffice --headless --convert-to 'doc:MS Word 97' breaks.docx
"""

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_BREAK

doc = Document()
p = doc.add_paragraph("Before the break")
p.add_run().add_break(WD_BREAK.PAGE)
p.add_run("after the break, same paragraph.")
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
doc.add_paragraph("Last paragraph of section one.")
doc.add_section(WD_SECTION.NEW_PAGE)
doc.add_paragraph("First paragraph of section two.")
doc.save("breaks.docx")
