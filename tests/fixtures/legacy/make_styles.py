# SPDX-FileCopyrightText: 2026 Jandira Technologies, LLC
# SPDX-License-Identifier: AGPL-3.0-only
"""styles.docx: a Title, Heading 1, Heading 8 and Heading 9 paragraph, and
an IF field whose code holds a nested PAGE field. LibreOffice's "MS Word
97" export of it is styles.doc:

    uv run --with python-docx python3 make_styles.py
    soffice --headless --convert-to 'doc:MS Word 97' styles.docx
"""

from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = Document()
doc.add_paragraph("Master Agreement", style="Title")
doc.add_paragraph("Scope", style="Heading 1")
doc.add_paragraph("Deep clause", style="Heading 8")
doc.add_paragraph("Deeper clause", style="Heading 9")
p = doc.add_paragraph("Page check: ")


def run(xml):
    p._p.append(parse_xml(f"<w:r {nsdecls('w')}>{xml}</w:r>"))


run('<w:fldChar w:fldCharType="begin"/>')
run('<w:instrText xml:space="preserve"> IF </w:instrText>')
run('<w:fldChar w:fldCharType="begin"/>')
run('<w:instrText xml:space="preserve"> PAGE </w:instrText>')
run('<w:fldChar w:fldCharType="separate"/>')
run("<w:t>1</w:t>")
run('<w:fldChar w:fldCharType="end"/>')
run('<w:instrText xml:space="preserve"> = 1 "first" "later" </w:instrText>')
run('<w:fldChar w:fldCharType="separate"/>')
run("<w:t>first</w:t>")
run('<w:fldChar w:fldCharType="end"/>')
doc.add_paragraph("Closing paragraph.")
doc.save("styles.docx")
